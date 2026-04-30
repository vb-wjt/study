"""Clone all unique (repo, branch) pairs from PI3.0-code.xlsx into D:\\cursor_workspace\\pi_origin.

- Layout: flat (one directory per repo, named by repo basename)
- Depth : --single-branch (target branch only, full history)
- Auth  : whatever git is already configured to use; GIT_TERMINAL_PROMPT=0 prevents hangs.
"""
from __future__ import annotations

import os
import re
import subprocess
import sys
import time
from pathlib import Path

import openpyxl

EXCEL = Path(r"C:\Users\Wu.juntao\Desktop\PI\PI next release\PI3.0-code.xlsx")
DEST_ROOT = Path(r"D:\cursor_workspace\pi_origin")
LOG_FILE = Path(r"D:\cursor_workspace\refactor_pi\clone_pi_repos.log")

assert DEST_ROOT.exists() and DEST_ROOT.is_dir(), f"Destination missing: {DEST_ROOT}"


def normalize_clone_url(raw: str) -> str:
    """Strip GitLab tree-view suffix and ensure .git suffix."""
    u = raw.strip()
    u = re.sub(r"/-/tree/.*$", "", u)
    if not u.endswith(".git"):
        u = u.rstrip("/") + ".git"
    return u


def collect_pairs() -> list[tuple[str, str]]:
    """Return list of unique (clone_url, branch) preserving first-seen order."""
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    seen: dict[tuple[str, str], None] = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = rows[0]
        has_header = any(
            isinstance(c, str) and c.lower() in ("repo", "classification", "branch")
            for c in header if c
        )
        start = 1 if has_header else 0
        for row in rows[start:]:
            url = next(
                (c for c in row if isinstance(c, str)
                 and c.startswith(("http://", "https://", "git@"))),
                None,
            )
            if not url:
                continue
            cells = [c for c in row if c]
            branch_cell = cells[-1] if len(cells) >= 2 else None
            branch = str(branch_cell).strip() if branch_cell and branch_cell != url else "develop"
            key = (normalize_clone_url(url), branch)
            seen.setdefault(key, None)
    return list(seen.keys())


def clone_one(clone_url: str, branch: str) -> tuple[str, float, str]:
    """Clone (or skip) a single repo. Returns (status, elapsed_seconds, message)."""
    repo_name = clone_url.rsplit("/", 1)[-1].removesuffix(".git")
    dest = DEST_ROOT / repo_name
    if dest.exists():
        return ("SKIP", 0.0, f"already exists: {dest}")

    cmd = [
        "git", "clone",
        "--single-branch",
        "--branch", branch,
        clone_url,
        str(dest),
    ]
    env = {**os.environ, "GIT_TERMINAL_PROMPT": "0"}
    t0 = time.time()
    proc = subprocess.run(cmd, env=env, capture_output=True, text=True, timeout=600)
    elapsed = time.time() - t0
    if proc.returncode == 0:
        return ("OK", elapsed, "")
    err = (proc.stderr or proc.stdout or "").strip().replace("\n", " | ")
    return ("FAIL", elapsed, err[:500])


def main() -> int:
    pairs = collect_pairs()
    print(f"Total unique repos to clone: {len(pairs)}")
    print(f"Destination: {DEST_ROOT}")
    print(f"Log file:    {LOG_FILE}")
    print()

    ok = fail = skip = 0
    with LOG_FILE.open("w", encoding="utf-8") as log:
        log.write(f"# Clone run started at {time.strftime('%Y-%m-%d %H:%M:%S')}\n")
        log.write(f"# Total: {len(pairs)}\n\n")
        for i, (url, branch) in enumerate(pairs, 1):
            name = url.rsplit("/", 1)[-1].removesuffix(".git")
            print(f"[{i:>3}/{len(pairs)}] {name:55s} (branch={branch})", flush=True)
            try:
                status, elapsed, msg = clone_one(url, branch)
            except subprocess.TimeoutExpired:
                status, elapsed, msg = ("FAIL", 600.0, "timeout after 600s")
            line = f"[{i:>3}/{len(pairs)}] {status:4s} {elapsed:6.1f}s  {name:55s} branch={branch:30s}"
            if msg:
                line += f"  -- {msg}"
            print("    " + line, flush=True)
            log.write(line + "\n")
            log.flush()
            if status == "OK":
                ok += 1
            elif status == "SKIP":
                skip += 1
            else:
                fail += 1

        summary = f"\n# Done. ok={ok} skip={skip} fail={fail}\n"
        log.write(summary)
        print(summary)
    return 0 if fail == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
